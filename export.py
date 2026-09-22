"""
export.py — Generates an Excel (.xlsx) export of the full sample list.
Uses openpyxl. File is generated fresh on each request.
"""

import tempfile
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from database import get_all_samples


def generate_export():
    """
    Build an .xlsx file with all non-deleted samples.
    Columns: Sample ID, Sample Name, Received From, Date Received, Status.
    Returns the file path to the generated spreadsheet.
    """
    samples = get_all_samples()

    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"

    # -- Header styling --
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2E4057', end_color='2E4057', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # -- Headers --
    headers = ['Sample ID', 'Sample Name', 'Received From', 'Date Received', 'Status']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # -- Data rows --
    for row_idx, sample in enumerate(samples, start=2):
        ws.cell(row=row_idx, column=1, value=sample['sample_id']).border = thin_border
        ws.cell(row=row_idx, column=2, value=sample['name']).border = thin_border
        ws.cell(row=row_idx, column=3, value=sample['received_from']).border = thin_border
        ws.cell(row=row_idx, column=4, value=sample['date_received']).border = thin_border
        ws.cell(row=row_idx, column=5, value=sample['status']).border = thin_border

    # -- Column widths --
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 18

    # -- Save to temp file --
    temp_dir = tempfile.mkdtemp()
    filepath = os.path.join(temp_dir, 'samples_export.xlsx')
    wb.save(filepath)

    return filepath
